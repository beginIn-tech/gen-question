"""
Text extraction service for various file formats
"""

import os
import io
from typing import List
import PyPDF2
from docx import Document
from fastapi import HTTPException, UploadFile


class TextExtractionService:
    """Service for extracting raw text from various file formats"""
    
    @staticmethod
    def get_supported_file_types() -> List[str]:
        """Return list of supported file extensions"""
        return ['.txt', '.pdf', '.doc', '.docx']
    
    @staticmethod
    def is_supported_file_type(filename: str) -> bool:
        """Check if file type is supported"""
        if not filename:
            return False
        
        file_extension = os.path.splitext(filename)[1].lower()
        return file_extension in TextExtractionService.get_supported_file_types()
    
    @staticmethod
    def _extract_from_txt(file_content: bytes) -> str:
        """Extract text from TXT file"""
        try:
            # Try UTF-8 first, then fall back to other encodings
            try:
                return file_content.decode('utf-8')
            except UnicodeDecodeError:
                try:
                    return file_content.decode('latin-1')
                except UnicodeDecodeError:
                    return file_content.decode('cp1252', errors='replace')
        except Exception as e:
            # Global error handling
            raise HTTPException(status_code=500, detail="Error extracting text from file")
    
    @staticmethod
    def _extract_from_pdf(file_content: bytes) -> str:
        """Extract text from PDF file"""
        try:
            text = ""
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            # Extract text from all pages
            for page_num in range(len(pdf_reader.pages)):
                page = pdf_reader.pages[page_num]
                text += page.extract_text() + "\n"
            
            if not text.strip():
                raise Exception("No text found in PDF")
            
            return text
            
        except Exception as e:
            # Global error handling
            raise HTTPException(status_code=500, detail="Error extracting text from PDF file")
    
    @staticmethod
    def _extract_from_docx(file_content: bytes) -> str:
        """Extract text from DOCX file (and DOC files saved as DOCX)"""
        try:
            doc_file = io.BytesIO(file_content)
            doc = Document(doc_file)
            
            text = ""
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            if not text.strip():
                raise Exception("No text found in document")
            
            return text
            
        except Exception as e:
            # Global error handling
            raise HTTPException(status_code=500, detail="Error extracting text from document file")
    
    @staticmethod
    async def extract_text_from_file(file: UploadFile) -> str:
        """
        Extract text from uploaded file based on file extension
        
        Args:
            file: The uploaded file object
            
        Returns:
            str: Extracted text content
            
        Raises:
            HTTPException: If file type is not supported or parsing fails
        """
        try:
            filename = file.filename
            
            if not filename:
                raise Exception("Filename is required")
            
            file_extension = os.path.splitext(filename)[1].lower()
            
            # Read file content
            file_content = await file.read()
            
            # Reset file pointer
            await file.seek(0)
            
            # Parse based on file type
            if file_extension == '.txt':
                return TextExtractionService._extract_from_txt(file_content)
            elif file_extension == '.pdf':
                return TextExtractionService._extract_from_pdf(file_content)
            elif file_extension in ['.doc', '.docx']:
                return TextExtractionService._extract_from_docx(file_content)
            else:
                supported_formats = TextExtractionService.get_supported_file_types()
                raise Exception(f"Unsupported file format: {file_extension}")
                
        except Exception as e:
            # Global error handling
            raise HTTPException(status_code=500, detail="Error extracting text from file")
